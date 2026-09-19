import os
import sys
import traceback

# Forçar o diretório de trabalho a ser a pasta onde o script ou executável está localizado
try:
    if getattr(sys, 'frozen', False):
        script_dir = os.path.dirname(os.path.abspath(sys.executable))
    else:
        script_dir = os.path.dirname(os.path.abspath(__file__))
    os.chdir(script_dir)
except Exception:
    pass

def registrar_erro_fatal(mensagem):
    try:
        with open("erro_analisador.txt", "w", encoding="utf-8") as f:
            f.write(mensagem)
    except Exception:
        pass

# Tentar realizar as importações e instalações necessárias de forma protegida
try:
    import subprocess
    from datetime import datetime

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        print("Instalando biblioteca openpyxl para o Excel...")
        try:
            subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except Exception as e:
            raise ImportError(f"Falha ao instalar openpyxl via pip: {str(e)}")

    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
    except ImportError:
        print("Instalando biblioteca reportlab para geracao de PDF...")
        try:
            subprocess.check_call([sys.executable, "-m", "pip", "install", "reportlab"])
            from reportlab.lib.pagesizes import A4
            from reportlab.lib import colors
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import cm
        except Exception as e:
            raise ImportError(f"Falha ao instalar reportlab via pip: {str(e)}")

    import tkinter as tk
    from tkinter import ttk, messagebox, filedialog
except Exception as e:
    info_erro = traceback.format_exc()
    registrar_erro_fatal(f"Erro nas dependencias de inicializacao:\n{info_erro}")
    print(f"Erro de dependencias:\n{info_erro}")
    sys.exit(1)

class AnalisadorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Analisador de Ocorrências Excel")
        self.root.geometry("540x630")
        self.root.configure(bg="#F5F6FA")
        
        # Centralizar a janela na tela
        window_width = 540
        window_height = 630
        screen_width = root.winfo_screenwidth()
        screen_height = root.winfo_screenheight()
        position_top = int(screen_height/2 - window_height/2)
        position_right = int(screen_width/2 - window_width/2)
        self.root.geometry(f'{window_width}x{window_height}+{position_right}+{position_top}')
        
        # Variáveis de estado para exportação em PDF
        self.dados_analisados = []
        self.meta_periodo_de = ""
        self.meta_periodo_ate = ""
        self.meta_arquivo = ""
        
        # Configurar Estilos Visuais ttk
        self.style = ttk.Style()
        self.style.theme_use("clam")
        
        self.style.configure("TLabel", font=("Segoe UI", 10), background="#F5F6FA", foreground="#2F3542")
        self.style.configure("TButton", font=("Segoe UI", 10, "bold"), background="#2F3640", foreground="#FFFFFF")
        self.style.map("TButton", background=[("active", "#353B48"), ("disabled", "#BDC3C7")], foreground=[("disabled", "#7F8C8D")])
        self.style.configure("Header.TLabel", font=("Segoe UI", 13, "bold"), background="#2F3640", foreground="#FFFFFF")
        
        # Banner de Cabeçalho
        self.header_frame = tk.Frame(root, bg="#2F3640", height=60)
        self.header_frame.pack(fill="x", side="top")
        self.header_label = ttk.Label(self.header_frame, text="ANALISADOR DE OCORRÊNCIAS", style="Header.TLabel")
        self.header_label.pack(pady=15)
        
        self.content_frame = tk.Frame(root, bg="#F5F6FA")
        self.content_frame.pack(fill="both", expand=True, padx=20, pady=15)
        
        # 1. Seleção do Arquivo
        ttk.Label(self.content_frame, text="Selecione o arquivo Excel:").pack(anchor="w", pady=3)
        self.file_frame = tk.Frame(self.content_frame, bg="#F5F6FA")
        self.file_frame.pack(fill="x", pady=3)
        
        self.file_var = tk.StringVar()
        self.file_combo = ttk.Combobox(self.file_frame, textvariable=self.file_var, state="readonly", font=("Segoe UI", 10))
        self.file_combo.pack(side="left", fill="x", expand=True)
        
        self.btn_browse = ttk.Button(self.file_frame, text="Buscar...", command=self.buscar_arquivo)
        self.btn_browse.pack(side="right", padx=(5, 0))
        
        self.atualizar_lista_arquivos()
        
        # Botão para Gerar Planilha Modelo
        self.btn_modelo = ttk.Button(self.content_frame, text="Criar Nova Planilha Modelo (.xlsx)", command=self.gerar_modelo)
        self.btn_modelo.pack(fill="x", pady=(5, 10))
        
        # 2. Filtros de Data
        self.dates_frame = tk.LabelFrame(self.content_frame, text=" Período de Análise ", bg="#F5F6FA", fg="#2F3542", font=("Segoe UI", 10, "bold"), padx=10, pady=10)
        self.dates_frame.pack(fill="x", pady=5)
        self.dates_frame.columnconfigure(1, weight=1)
        
        ttk.Label(self.dates_frame, text="Data Inicial (De):").grid(row=0, column=0, sticky="w", pady=5)
        self.entry_de = ttk.Entry(self.dates_frame, font=("Segoe UI", 10))
        self.entry_de.grid(row=0, column=1, sticky="ew", padx=(10, 0), pady=5)
        self.entry_de.insert(0, "01/07/2026")
        
        ttk.Label(self.dates_frame, text="Data Final (Até):").grid(row=1, column=0, sticky="w", pady=5)
        self.entry_ate = ttk.Entry(self.dates_frame, font=("Segoe UI", 10))
        self.entry_ate.grid(row=1, column=1, sticky="ew", padx=(10, 0), pady=5)
        self.entry_ate.insert(0, "31/12/2026")
        
        # 3. Botões de Ação (Analisar e Gerar PDF)
        self.actions_frame = tk.Frame(self.content_frame, bg="#F5F6FA")
        self.actions_frame.pack(fill="x", pady=10)
        self.actions_frame.columnconfigure(0, weight=1)
        self.actions_frame.columnconfigure(1, weight=1)
        
        self.btn_analisar = ttk.Button(self.actions_frame, text="🔍 Analisar Ocorrências", command=self.analisar)
        self.btn_analisar.grid(row=0, column=0, sticky="ew", padx=(0, 4))
        
        self.btn_pdf = ttk.Button(self.actions_frame, text="📄 Gerar Relatório PDF", command=self.gerar_pdf, state="disabled")
        self.btn_pdf.grid(row=0, column=1, sticky="ew", padx=(4, 0))
        
        # 4. Tabela de Resultados
        ttk.Label(self.content_frame, text="Resultado da Análise (Ordenado pelas repetições):").pack(anchor="w", pady=3)
        
        self.tree_frame = tk.Frame(self.content_frame)
        self.tree_frame.pack(fill="both", expand=True)
        
        self.scrollbar = ttk.Scrollbar(self.tree_frame)
        self.scrollbar.pack(side="right", fill="y")
        
        self.tree = ttk.Treeview(self.tree_frame, columns=("numero", "repeticoes"), show="headings", yscrollcommand=self.scrollbar.set)
        self.tree.heading("numero", text="Número")
        self.tree.heading("repeticoes", text="Repetições")
        self.tree.column("numero", anchor="center", width=150)
        self.tree.column("repeticoes", anchor="center", width=150)
        self.tree.pack(fill="both", expand=True)
        
        self.scrollbar.config(command=self.tree.yview)
        
    def buscar_arquivo(self):
        filepath = filedialog.askopenfilename(filetypes=[("Arquivos Excel", "*.xlsx")])
        if filepath:
            self.file_var.set(filepath)
            
    def atualizar_lista_arquivos(self):
        try:
            arquivos = [f for f in os.listdir(".") if f.endswith(".xlsx")]
            self.file_combo['values'] = arquivos
            if arquivos:
                if "Analise_Ocorrencias.xlsx" in arquivos:
                    idx = arquivos.index("Analise_Ocorrencias.xlsx")
                    self.file_combo.current(idx)
                else:
                    self.file_combo.current(0)
        except Exception:
            pass
            
    def gerar_modelo(self):
        filepath = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Arquivos Excel", "*.xlsx")],
            initialfile="Modelo_Lancamentos.xlsx",
            title="Salvar Planilha Modelo"
        )
        if not filepath:
            return
            
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Lançamentos"
            ws.views.sheetView[0].showGridLines = True
            
            # Estilizar cabeçalhos (Tema moderno Slate) usando classes importadas
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="2F3640", end_color="2F3640", fill_type="solid")
            center_align = Alignment(horizontal="center", vertical="center")
            border_thin = Side(border_style="thin", color="DCDDE1")
            border_data = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
            
            ws["A1"] = "Número"
            ws["B1"] = "Data"
            
            for col in ["A1", "B1"]:
                ws[col].font = header_font
                ws[col].fill = header_fill
                ws[col].alignment = center_align
                ws[col].border = border_data
            
            # Formatação inicial de célula vazia
            ws.cell(row=2, column=1).font = Font(name="Segoe UI", size=11)
            ws.cell(row=2, column=1).alignment = center_align
            ws.cell(row=2, column=1).border = border_data
            
            cell_date = ws.cell(row=2, column=2)
            cell_date.font = Font(name="Segoe UI", size=11)
            cell_date.alignment = center_align
            cell_date.number_format = 'dd/mm/yyyy'
            cell_date.border = border_data
            
            ws.column_dimensions['A'].width = 18
            ws.column_dimensions['B'].width = 18
            
            wb.save(filepath)
            messagebox.showinfo("Sucesso", f"Planilha modelo criada com sucesso!\nSalva em: {filepath}")
            
            self.atualizar_lista_arquivos()
            
            filename = os.path.basename(filepath)
            values = self.file_combo['values']
            if filename in values:
                idx = values.index(filename)
                self.file_combo.current(idx)
            else:
                self.file_var.set(filepath)
                
        except Exception as e:
            messagebox.showerror("Erro ao criar modelo", f"Não foi possível salvar o modelo:\n{str(e)}")
            
    def analisar(self):
        filepath = self.file_var.get()
        if not filepath:
            messagebox.showerror("Erro", "Por favor, selecione um arquivo Excel.")
            return
            
        de_str = self.entry_de.get().strip()
        ate_str = self.entry_ate.get().strip()
        
        try:
            date_de = datetime.strptime(de_str, "%d/%m/%Y")
            date_ate = datetime.strptime(ate_str, "%d/%m/%Y")
        except ValueError:
            messagebox.showerror("Erro de Formato", "As datas devem estar no formato DD/MM/AAAA.\nExemplo: 10/07/2026")
            return
            
        if date_de > date_ate:
            messagebox.showerror("Erro de Intervalo", "A data inicial não pode ser maior que a data final.")
            return
            
        for item in self.tree.get_children():
            self.tree.delete(item)
            
        self.dados_analisados = []
        self.btn_pdf.config(state="disabled")
            
        try:
            wb = openpyxl.load_workbook(filepath, data_only=True)
            ws = wb.active
            
            registros_vistos = set()
            contagem = {}
            
            for row in range(2, ws.max_row + 1):
                cell_num = ws.cell(row=row, column=1).value
                cell_date = ws.cell(row=row, column=2).value
                
                if cell_num is None:
                    continue
                
                date_val = None
                if isinstance(cell_date, datetime):
                    date_val = cell_date
                elif cell_date:
                    cleaned_date = str(cell_date).strip()
                    if len(cleaned_date) == 9 and cleaned_date[2] == '/':
                        cleaned_date = cleaned_date[:5] + '/' + cleaned_date[5:]
                    
                    for fmt in ("%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
                        try:
                            date_val = datetime.strptime(cleaned_date, fmt)
                            break
                        except ValueError:
                            continue
                
                if not date_val:
                    continue
                
                if date_de <= date_val <= date_ate:
                    num_key = str(cell_num)
                    date_str = date_val.strftime("%Y-%m-%d")
                    par_unico = (num_key, date_str)
                    
                    if par_unico not in registros_vistos:
                        registros_vistos.add(par_unico)
                        contagem[num_key] = contagem.get(num_key, 0) + 1
            
            resultados_ordenados = sorted(contagem.items(), key=lambda x: x[1], reverse=True)
            
            if not resultados_ordenados:
                messagebox.showinfo("Informação", "Nenhum número foi encontrado dentro deste período de datas.")
                return
                
            self.dados_analisados = resultados_ordenados
            self.meta_periodo_de = de_str
            self.meta_periodo_ate = ate_str
            self.meta_arquivo = os.path.basename(filepath)
            self.btn_pdf.config(state="normal")
                
            for num, reps in resultados_ordenados:
                self.tree.insert("", "end", values=(num, reps))
                
        except Exception as e:
            self.dados_analisados = []
            self.btn_pdf.config(state="disabled")
            messagebox.showerror("Erro ao ler arquivo", f"Não foi possível processar o arquivo Excel:\n{str(e)}")

    def gerar_pdf(self):
        if not self.dados_analisados:
            messagebox.showwarning("Aviso", "Não há dados analisados para gerar o relatório PDF.")
            return

        data_sugestao = datetime.now().strftime("%d-%m-%Y")
        sugestao_nome = f"Relatorio_Ocorrencias_{data_sugestao}.pdf"
        
        filepath = filedialog.asksaveasfilename(
            defaultextension=".pdf",
            filetypes=[("Documentos PDF", "*.pdf")],
            initialfile=sugestao_nome,
            title="Salvar Relatório PDF"
        )
        if not filepath:
            return

        try:
            # Configuração do Documento A4 com margens proporcionais
            doc = SimpleDocTemplate(
                filepath,
                pagesize=A4,
                leftMargin=1.8 * cm,
                rightMargin=1.8 * cm,
                topMargin=1.8 * cm,
                bottomMargin=1.8 * cm
            )

            styles = getSampleStyleSheet()
            
            # Estilos Customizados
            style_titulo = ParagraphStyle(
                'TituloRelatorio',
                parent=styles['Heading1'],
                fontName='Helvetica-Bold',
                fontSize=18,
                leading=22,
                textColor=colors.HexColor("#2F3640"),
                alignment=1,
                spaceAfter=4
            )

            style_subtitulo = ParagraphStyle(
                'SubtituloRelatorio',
                parent=styles['Normal'],
                fontName='Helvetica',
                fontSize=10,
                leading=14,
                textColor=colors.HexColor("#718093"),
                alignment=1,
                spaceAfter=10
            )

            style_meta_label = ParagraphStyle(
                'MetaLabel',
                parent=styles['Normal'],
                fontName='Helvetica-Bold',
                fontSize=9,
                leading=12,
                textColor=colors.HexColor("#2F3640")
            )

            style_meta_val = ParagraphStyle(
                'MetaVal',
                parent=styles['Normal'],
                fontName='Helvetica',
                fontSize=9,
                leading=12,
                textColor=colors.HexColor("#353B48")
            )

            style_th = ParagraphStyle(
                'TableHeader',
                parent=styles['Normal'],
                fontName='Helvetica-Bold',
                fontSize=10,
                leading=13,
                textColor=colors.white,
                alignment=1
            )

            style_td = ParagraphStyle(
                'TableCell',
                parent=styles['Normal'],
                fontName='Helvetica',
                fontSize=9,
                leading=12,
                textColor=colors.HexColor("#2F3542"),
                alignment=1
            )

            style_td_bold = ParagraphStyle(
                'TableCellBold',
                parent=styles['Normal'],
                fontName='Helvetica-Bold',
                fontSize=9,
                leading=12,
                textColor=colors.HexColor("#2F3640"),
                alignment=1
            )

            elementos = []

            # Cabeçalho Principal
            elementos.append(Paragraph("RELATÓRIO DE OCORRÊNCIAS DE HINOS", style_titulo))
            elementos.append(Paragraph("Consolidação estatística e ranking de repetições por período", style_subtitulo))
            elementos.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#2F3640"), spaceAfter=12))

            # Caixa de Metadados / Resumo da Consulta
            agora_str = datetime.now().strftime("%d/%m/%Y às %H:%M:%S")
            total_hinos = len(self.dados_analisados)
            total_presencas = sum(reps for _, reps in self.dados_analisados)

            dados_meta = [
                [
                    Paragraph("<b>Arquivo de Origem:</b>", style_meta_label),
                    Paragraph(self.meta_arquivo, style_meta_val),
                    Paragraph("<b>Data de Emissão:</b>", style_meta_label),
                    Paragraph(agora_str, style_meta_val)
                ],
                [
                    Paragraph("<b>Período Analisado:</b>", style_meta_label),
                    Paragraph(f"{self.meta_periodo_de} até {self.meta_periodo_ate}", style_meta_val),
                    Paragraph("<b>Total de Hinos:</b>", style_meta_label),
                    Paragraph(f"{total_hinos} hinos ({total_presencas} execuções)", style_meta_val)
                ]
            ]

            tabela_meta = Table(dados_meta, colWidths=[4.0 * cm, 4.7 * cm, 4.0 * cm, 4.7 * cm])
            tabela_meta.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F5F6FA")),
                ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#DCDDE1")),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('TOPPADDING', (0, 0), (-1, -1), 6),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                ('LEFTPADDING', (0, 0), (-1, -1), 8),
                ('RIGHTPADDING', (0, 0), (-1, -1), 8),
            ]))
            elementos.append(tabela_meta)
            elementos.append(Spacer(1, 14))

            # Tabela de Resultados
            tabela_dados = [
                [
                    Paragraph("Posição", style_th),
                    Paragraph("Número do Hino", style_th),
                    Paragraph("Total de Repetições", style_th)
                ]
            ]

            for idx, (num, reps) in enumerate(self.dados_analisados, 1):
                tabela_dados.append([
                    Paragraph(f"{idx}º", style_td_bold),
                    Paragraph(str(num), style_td),
                    Paragraph(str(reps), style_td_bold)
                ])

            tabela_resultado = Table(tabela_dados, colWidths=[3.5 * cm, 7.5 * cm, 6.4 * cm], repeatRows=1)
            
            estilo_tabela = [
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#2F3640")),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('TOPPADDING', (0, 0), (-1, -1), 5),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#DCDDE1")),
            ]

            # Cores zebradas para linhas de dados
            for i in range(1, len(tabela_dados)):
                if i % 2 == 0:
                    estilo_tabela.append(('BACKGROUND', (0, i), (-1, i), colors.HexColor("#F8F9FA")))
                else:
                    estilo_tabela.append(('BACKGROUND', (0, i), (-1, i), colors.white))

            tabela_resultado.setStyle(TableStyle(estilo_tabela))
            elementos.append(tabela_resultado)
            
            elementos.append(Spacer(1, 12))
            elementos.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#DCDDE1"), spaceAfter=6))
            
            style_rodape = ParagraphStyle(
                'RodapeRelatorio',
                parent=styles['Normal'],
                fontName='Helvetica-Oblique',
                fontSize=8,
                leading=10,
                textColor=colors.HexColor("#718093"),
                alignment=1
            )
            elementos.append(Paragraph("Relatório gerado automaticamente pelo Analisador de Ocorrências Excel", style_rodape))

            # Construir o arquivo PDF
            doc.build(elementos)

            abrir = messagebox.askyesno(
                "Sucesso",
                f"Relatório PDF gerado com sucesso!\n\nSalvo em:\n{filepath}\n\nDeseja abrir o arquivo agora?"
            )
            if abrir:
                try:
                    os.startfile(filepath)
                except Exception as e_open:
                    messagebox.showwarning("Aviso", f"PDF gerado, mas não foi possível abrir automaticamente:\n{str(e_open)}")

        except Exception as e:
            messagebox.showerror("Erro ao Gerar PDF", f"Não foi possível salvar o relatório PDF:\n{str(e)}")

if __name__ == "__main__":
    try:
        root = tk.Tk()
        app = AnalisadorApp(root)
        root.mainloop()
    except Exception as e:
        info_erro = traceback.format_exc()
        registrar_erro_fatal(f"Erro durante a execucao do aplicativo:\n{info_erro}")
        print(f"Erro fatal de execucao:\n{info_erro}")
        sys.exit(1)
